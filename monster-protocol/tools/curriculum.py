"""Build editable classroom handouts and a small portable curriculum ZIP.
Run with python3 tools/curriculum.py; requires python-docx 1.2.0 (tools/requirements-curriculum.txt).
This script never edits or regenerates game projects.
"""
from pathlib import Path
import re, zipfile, hashlib
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'downloads'
BASE='https://github.com/olgaschev418-tech/my-first-project/blob/main/monster-protocol/docs/'


def inline(paragraph, text):
    pattern=r'(\[[^\]]+\]\([^)]+\)|\*\*[^*]+\*\*|`[^`]+`)'
    for part in re.split(pattern,text):
        if part.startswith('[') and '](' in part:
            label,target=re.match(r'\[([^\]]+)\]\(([^)]+)\)',part).groups()
            url=target if target.startswith(('https://','http://')) else BASE+target
            link=OxmlElement('w:hyperlink');link.set(qn('r:id'),paragraph.part.relate_to(url,RT.HYPERLINK,is_external=True))
            run=OxmlElement('w:r');props=OxmlElement('w:rPr');color=OxmlElement('w:color');color.set(qn('w:val'),'245A81');props.append(color);run.append(props)
            value=OxmlElement('w:t');value.text=label;run.append(value);link.append(run);paragraph._p.append(link)
        else:
            bold=part.startswith('**') and part.endswith('**')
            code=part.startswith('`') and part.endswith('`')
            run=paragraph.add_run(part[2:-2] if bold else part[1:-1] if code else part)
            run.bold=bold
            if code:run.font.name='Consolas';run.font.size=Pt(9)


def make_doc(source,target,subtitle,page_cards=False):
    doc=Document();doc.core_properties.title=subtitle;doc.core_properties.author='AI GAME';doc.core_properties.subject='GDevelop: 8 занятий по 90 минут, возраст 8–10 лет'
    section=doc.sections[0]
    section.page_width=Cm(21);section.page_height=Cm(29.7)
    section.top_margin=Cm(1.7);section.bottom_margin=Cm(1.7);section.left_margin=Cm(1.8);section.right_margin=Cm(1.8)
    normal=doc.styles['Normal'];normal.font.name='DejaVu Sans';normal.font.size=Pt(10)
    normal.paragraph_format.space_after=Pt(6)
    lang=OxmlElement('w:lang');lang.set(qn('w:val'),'ru-RU');normal.element.get_or_add_rPr().append(lang)
    for name in ['Heading 1','Heading 2','Heading 3']:
        doc.styles[name].font.name='DejaVu Sans';doc.styles[name].font.color.rgb=RGBColor.from_string('245A60')
    section.header.paragraphs[0].text='AI GAME · GDevelop · 8–10 лет · 8 × 90 минут'
    footer=section.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run('Учебные материалы · ')
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
    lines=(ROOT/'docs'/source).read_text(encoding='utf-8').splitlines();i=0;card_count=0
    while i<len(lines):
        line=lines[i].strip()
        if not line:i+=1;continue
        if line.startswith('|'):
            rows=[]
            while i<len(lines) and lines[i].strip().startswith('|'):
                cells=[c.strip() for c in lines[i].strip().strip('|').split('|')]
                if not all(re.fullmatch(r':?-+:?',c) for c in cells):rows.append(cells)
                i+=1
            table=doc.add_table(rows=0,cols=max(map(len,rows)));table.style='Light Shading Accent 1'
            for ri,row in enumerate(rows):
                cells=table.add_row().cells
                for cell,text in zip(cells,row):
                    inline(cell.paragraphs[0],text)
                    for run in cell.paragraphs[0].runs:run.font.size=Pt(9);run.bold=ri==0 or run.bold
                trpr=cells[0]._tc.getparent().get_or_add_trPr();trpr.append(OxmlElement('w:cantSplit'))
            if rows[0][0] in ('Минута','Время'):
                widths=[1.3,2.4,7.4,6.3] if len(rows[0])==4 else [1.4,10,6] if len(rows[0])==3 else None
                if widths:
                    table.autofit=False
                    for column,width in zip(table.columns,widths):column.width=Cm(width)
                    for row in table.rows:
                        for cell,width in zip(row.cells,widths):cell.width=Cm(width)
            table.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
            doc.add_paragraph();continue
        match=re.match(r'^(#{1,4})\s+(.*)',line)
        if match:
            depth=len(match[1]);title=match[2]
            if depth==1:
                doc.add_heading(title,0);doc.add_paragraph(subtitle,'Subtitle')
            else:
                p=doc.add_heading(title,min(depth-1,3))
                if depth==2 and re.match(r'^\d+\.',title):
                    if page_cards and card_count>0:p.paragraph_format.page_break_before=True
                    if source=='LESSONS_RU.md':p.paragraph_format.page_break_before=True
                    card_count+=1
            if source=='CLICK_BY_CLICK_RU.md' and depth==2 and title.startswith('Занятие '):p.paragraph_format.page_break_before=True
            i+=1;continue
        if re.match(r'^- \[[ x]\] ',line):p=doc.add_paragraph();inline(p,'☐ '+line[6:])
        elif line.startswith('- '):p=doc.add_paragraph(style='List Bullet');inline(p,line[2:])
        elif re.match(r'^\d+\. ',line):p=doc.add_paragraph(style='List Number');inline(p,re.sub(r'^\d+\. ','',line))
        else:p=doc.add_paragraph();inline(p,line)
        i+=1
    doc.save(OUT/target)


def build():
    OUT.mkdir(exist_ok=True)
    specs=[
        ('FIRST_LESSON_CLICKS_RU.md','AI_GAME_Lesson_01_Click_by_Click_RU.docx','Первое занятие · все 90 минут · приложения, нажатия и проверки',False),
        ('CLICK_BY_CLICK_RU.md','AI_GAME_All_8_Lessons_Click_by_Click_RU.docx','Все восемь встреч · маршруты действий и сценарии преподавателя',False),
        ('LESSONS_RU.md','AI_GAME_8x90_Teacher_Plan_RU.docx','Подробный план преподавателя · два блока по четыре встречи',False),
        ('OFFSCREEN_GAMES_RU.md','AI_GAME_8x90_Activity_Cards_RU.docx','Восемь 15-минутных игр · правила и роли · карточки для печати',True),
        ('YEAR_PLAN_RU.md','AI_GAME_Year_Plan_RU.docx','Предложение модульной программы · 32 или 36 встреч',False),
        ('STUDENT_CARDS_RU.md','AI_GAME_8x90_Student_Passport_RU.docx','Личные задания · отметки миссий · портфолио',False),
    ]
    for source,target,subtitle,pages in specs:make_doc(source,target,subtitle,pages)
    readme='''# Учебный комплект AI GAME · 8 × 90 минут

- AI_GAME_Lesson_01_Click_by_Click_RU.docx — первое занятие: 90 строк по одной минуте, все нажатия.
- AI_GAME_All_8_Lessons_Click_by_Click_RU.docx — сценарии всех встреч и карты действий в редакторе.
- AI_GAME_8x90_Teacher_Plan_RU.docx — общий методический план преподавателя.
- AI_GAME_8x90_Activity_Cards_RU.docx — восемь игр без компьютера.
- AI_GAME_Year_Plan_RU.docx — идея модульной программы на год.
- AI_GAME_8x90_Student_Passport_RU.docx — карточки ученика.
- docs/ — те же тексты и технические памятки в Markdown.

Это только учебные материалы. Рабочую игру и ученический шаблон менять из-за нового расписания не нужно. Ссылки: https://github.com/olgaschev418-tech/my-first-project/tree/main/monster-protocol .

Возраст 8–10 лет. Два абонемента по четыре встречи. На встрече 4 — первая играбельная версия, на встрече 8 — итоговый проект. В каждом занятии есть 15 минут игры без компьютера и отдельный короткий перерыв. Будущие годовые модули, кроме «Мой монстр», пока являются предложением.
'''
    target=OUT/'AI_GAME_Curriculum_8x90_RU.zip'
    with zipfile.ZipFile(target,'w',zipfile.ZIP_DEFLATED,compresslevel=9) as z:
        prefix='AI_GAME_8x90_RU/'
        z.writestr(prefix+'README.md',readme)
        for _,name,_,_ in specs:z.write(OUT/name,prefix+name)
        for name in ['FIRST_LESSON_CLICKS_RU.md','CLICK_BY_CLICK_RU.md','LESSONS_RU.md','OFFSCREEN_GAMES_RU.md','YEAR_PLAN_RU.md','STUDENT_CARDS_RU.md','TEACHER_START_RU.md','EVENTS_RU.md','AI_RU.md']:
            z.write(ROOT/'docs'/name,prefix+'docs/'+name)
    with zipfile.ZipFile(target) as z:assert z.testzip() is None
    checksum=hashlib.sha256(target.read_bytes()).hexdigest()
    (OUT/'AI_GAME_Curriculum_8x90_RU.sha256').write_text(checksum+'  '+target.name+'\n')
    print(f'Built six Word documents and {target.name} ({target.stat().st_size:,} bytes); CRC OK.')


if __name__=='__main__':build()
