"""Сборка отчёта с точными абзацными параметрами Proga_otchet.docx."""
import argparse
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt
from docx.oxml.ns import qn
from PIL import Image

ROOT = Path(__file__).resolve().parents[1]
DEFAULT = ROOT / 'templates' / 'report_format.docx'
parser = argparse.ArgumentParser()
parser.add_argument('--template', type=Path, default=DEFAULT)
args = parser.parse_args()
doc = Document(args.template)
paragraphs = list(doc.paragraphs)
# These Word properties and spacer paragraphs are copied from Proga_otchet.docx.
roles = {p.text[1:]: deepcopy(p._p.pPr) for p in paragraphs
         if p.text.startswith('@')}

for index, content in [
        (14, 'ЛАБОРАТОРНАЯ РАБОТА №3'),
        (16, 'Тема: «МНОЖЕСТВЕННОЕ НАСЛЕДОВАНИЕ»'),
        (23, 'Направленность (профиль)\t Искусственный интеллект и '
             'математическое моделирование'),
        (28, 'Обучающаяся\t\t             _____________________       '
             'Шевченко О.Н.')]:
    p = paragraphs[index]
    formatting = deepcopy(p.runs[0]._r.rPr) if p.runs else None
    p.clear()
    r = p.add_run(content)
    if formatting is not None:
        r._r.insert(0, formatting)
for paragraph in paragraphs[43:]:
    paragraph._p.getparent().remove(paragraph._p)


def paragraph(role):
    p = doc.add_paragraph()
    p._p.insert(0, deepcopy(roles[role]))
    return p


def gap(role, keep=False):
    p = paragraph(role)
    if keep:
        p.paragraph_format.keep_with_next = True
    return p


def text(content, bold=False, kind='body'):
    p = paragraph(kind)
    r = p.add_run(content)
    r.font.name = 'Courier New' if kind == 'code' else 'Times New Roman'
    r.font.size = Pt(12 if kind == 'code' else 14)
    r.bold = bold
    if kind in ('body', 'conclusion'):
        p.paragraph_format.widow_control = True
    if kind == 'code':
        # Keep source indentation and the sample's single line spacing.
        if content.rstrip().endswith(':') or content.lstrip().startswith('@'):
            p.paragraph_format.keep_with_next = True
    return p


def heading(content, centered=False):
    if doc.paragraphs[-1].text.strip():
        gap('work_before' if centered else 'heading_before')
    p = paragraph('work_heading' if centered else 'heading')
    r = p.add_run(content)
    r.font.name = 'Times New Roman'
    r.font.size = Pt(14)
    r.bold = True
    p.paragraph_format.keep_with_next = True
    gap('work_after' if centered else 'heading_after', keep=True)
    return p


def figure(number, caption, width=16, height=19):
    path = ROOT / 'assets' / f'figure{number}.png'
    w, h = Image.open(path).size
    role = 'image' if number in (1, 3, 5, 6, 11) else 'diagram'
    p = paragraph(role)
    p.paragraph_format.keep_with_next = True
    r = p.add_run()
    r.add_picture(str(path), width=Cm(min(width, height * w / h)))
    r._r.xpath('.//wp:docPr')[0].set('descr', caption)
    text(f'Рисунок {number} – {caption}', kind='caption')


def listing(number, filename):
    gap('listing_before')
    p = text(f'Листинг {number} – Код программы', kind='listing')
    p.paragraph_format.keep_with_next = True
    for index, line in enumerate((ROOT / 'programs' / filename).read_text(
            encoding='utf-8').splitlines()):
        p = text(line, kind='code')
        if index < 4:
            p.paragraph_format.keep_with_next = True
    gap('listing_after')


p = paragraph('goal')
p.paragraph_format.page_break_before = True
p.add_run('Цель работы: ').bold = True
p.add_run('освоить принципы множественного наследования в Python, понять '
          'механизм MRO (Method Resolution Order), научиться проектировать '
          'классы с использованием нескольких родительских классов и '
          'разрешать потенциальные конфликты методов.')
for index, r in enumerate(p.runs):
    r.font.name = 'Times New Roman'
    r.font.size = Pt(14)
    r.bold = index == 0
heading('Задание')
text('1. Написать классы по описанию варианта 2 «Сотрудники компании». '
     'Базовый Employee содержит уникальный id. Классы ManagerSkills и '
     'DeveloperSkills реализуют work; Manager наследует Employee и '
     'ManagerSkills, Developer — Employee и DeveloperSkills. Создать общий '
     'список сотрудников и вызвать work в цикле. Объяснить возможную '
     'трудность множественного наследования и способы её разрешения.')
text('2. Для направления 09.03.02 «Информационные системы и технологии», '
     'профиля «Искусственный интеллект и математическое моделирование», '
     'выполнить вариант 2 «Слои нейросети»: Layer с полем id; DenseLayer и '
     'ConvolutionalLayer с методами forward и backward; классы Dense и Conv '
     'с множественным наследованием от Layer и соответствующего класса '
     'поведения.')
text('Выполнить также общие требования задания 2: разработать прототип '
     'информационно-аналитического ядра. Предусмотреть базовый класс с id '
     'и metadata; миксины валидации, кэширования, журналирования и экспорта; '
     'MLModel, EconomicIndicator и HybridAnalyzer. Продемонстрировать '
     'обучение, прогноз, экономическую оценку, super(), вывод __mro__, '
     'обработку исключений и проверку типов. Для профиля ИИ показать '
     'метрики качества и подбор гиперпараметра. Построить UML-диаграммы '
     'классов и последовательности сценария создания объекта, загрузки '
     'данных, обучения, расчёта индикатора и экспорта отчёта.')
heading('Теоретические сведения')
text('Множественное наследование позволяет классу иметь несколько '
     'родителей. Python определяет порядок поиска методов с помощью '
     'C3-линеаризации: он доступен через Class.__mro__ и Class.mro(). '
     'Порядок сохраняет приоритет родителей и согласованность уже '
     'построенных иерархий. Каждый класс встречается в MRO один раз.')
text('Миксин добавляет отдельную функциональность и используется вместе '
     'с основным классом. Он может не иметь состояния, как ValidationMixin, '
     'или хранить служебное состояние, как кэш и журнал. super() передаёт '
     'вызов следующему подходящему классу в MRO фактического объекта. '
     'Это особенно важно при ромбовидной иерархии: согласованные сигнатуры '
     'и однократный вызов super() предотвращают повторную инициализацию '
     'общего предка.')
heading('Ход работы', centered=True)

heading('1. Сотрудники компании')
text('Я реализовала Employee с общим счётчиком идентификаторов. '
     'ManagerSkills.work возвращает строку «id-й управляет», '
     'DeveloperSkills.work — «id-й программирует». Классы Manager и '
     'Developer объединяют состояние Employee с нужным поведением. '
     'Создала список из четырёх сотрудников и вызвала work для каждого '
     'элемента (см. рисунки 1, 2). Идентификаторы уникальны среди созданных '
     'сотрудников в пределах одного запуска программы.')
text('В исходных Manager и Developer конфликта work нет: Employee этот '
     'метод не определяет. Потенциальная неоднозначность возникает при '
     'объединении ManagerSkills и DeveloperSkills в одном классе. Без '
     'переопределения Python выбирает первый подходящий метод по MRO. '
     'Для TeamLead я явно переопределила work и вызвала оба метода '
     'родителей, объединив результаты. Такой прямой вызов подходит '
     'для независимых методов; для кооперативной цепочки инициализации '
     'в аналитическом ядре используется super().')
figure(1, 'Результат работы программы «Сотрудники компании»', height=15)
figure(2, 'UML-диаграмма классов сотрудников компании')
listing(1, 'task1_employees.py')

heading('2. Слои нейросети и информационно-аналитическое ядро')
heading('2.1. Слои нейросети — вариант 2')
text('Я реализовала Layer с уникальным id, миксины поведения DenseLayer '
     'и ConvolutionalLayer, а также классы Dense(Layer, DenseLayer) и '
     'Conv(Layer, ConvolutionalLayer). Для каждого экземпляра вызываются '
     'forward и backward, выводится MRO (см. рисунки 3, 4). Оба класса '
     'поведения предоставляют оба прохода. В соответствии с вариантом '
     'методы возвращают строки: численные операции с тензорами и '
     'обратное распространение градиента в этом учебном примере не '
     'реализуются.')
text('MRO для Dense: Dense → Layer → DenseLayer → object. Для Conv: '
     'Conv → Layer → ConvolutionalLayer → object. Конструктор находится '
     'в Layer, методы проходов — в соответствующем классе поведения. '
     'Таким образом, интерфейс объектов одинаков, а источник '
     'унаследованной реализации различается.')
figure(3, 'Результат работы программы «Слои нейросети»')
figure(4, 'UML-диаграмма классов Layer, Dense и Conv')
listing(2, 'task2_layers.py')

heading('2.2. Архитектура аналитического ядра')
text('Чтобы выполнить общую часть задания 2, я разработала отдельный '
     'модуль analytics_core.py. BaseComponent создаёт общий уникальный id '
     'и копирует metadata. ValidationMixin проверяет числовые данные; '
     'CachingMixin повторно использует прогнозы; LoggingMixin сохраняет '
     'операции в списке events; ReportMixin экспортирует словарь в '
     'JSON-строку и, по указанному пути, в UTF-8-файл. Служебное '
     'состояние кэша и журнала создаётся кооперативными конструкторами.')
text('MLModel реализует одномерную kNN-регрессию: хранит обучающие пары, '
     'выбирает k ближайших по абсолютному расстоянию значений x и '
     'возвращает среднее соответствующих y. При равных расстояниях '
     'выбирается более ранняя запись. Метод tune_k сравнивает кандидаты '
     'с помощью leave-one-out: каждая точка по очереди исключается из '
     'обучения и используется только для проверки. Критерий — минимум '
     'среднеквадратичной ошибки MSE; при равенстве выбирается меньшее k. '
     'После подбора модель обучается на всех данных с лучшим k.')
text('Метод metrics вычисляет MAE = Σ|ŷᵢ − yᵢ| / n и '
     'RMSE = √(Σ(ŷᵢ − yᵢ)² / n). В демонстрации эти две метрики '
     'вычисляются на обучающем наборе и характеризуют подгонку, '
     'а не качество на независимой тестовой выборке. При fit и '
     'load_data кэш очищается; после загрузки новых данных требуется '
     'повторное обучение.')
text('EconomicIndicator вычисляет ROI = (R − C) / C · 100 %, где R — '
     'выручка, C — положительные затраты. Положительный ROI '
     'интерпретируется как прибыль, отрицательный — как убыток, нулевой '
     '— как безубыточность. HybridAnalyzer наследует MLModel и '
     'EconomicIndicator, получает прогноз выручки и передаёт его в '
     'экономический расчёт. Это учебный сценарий оценки по прогнозу; '
     'модель kNN не экстраполирует тренд за пределы обучающего диапазона.')

heading('2.3. MRO, super() и конфликты методов')
text('В программе создаются отдельные экземпляры MLModel и '
     'EconomicIndicator, затем HybridAnalyzer. Фактический MRO '
     'гибридного класса: HybridAnalyzer → MLModel → EconomicIndicator '
     '→ ValidationMixin → CachingMixin → LoggingMixin → ReportMixin '
     '→ BaseComponent → object (см. рисунок 5). Один набор общих '
     'миксинов и BaseComponent включается в порядок поиска только один '
     'раз, несмотря на несколько путей наследования.')
text('В конструкторах сначала вызывается super().__init__, после '
     'возврата добавляется запись в init_trace. Поэтому список '
     'завершения инициализации идёт от общего предка к потомку: '
     'BaseComponent → LoggingMixin → CachingMixin → EconomicIndicator '
     '→ MLModel → HybridAnalyzer. Это порядок завершения '
     'кооперативных конструкторов, а не новая последовательность MRO.')
text('Метод describe переопределён в обоих основных родителях. '
     'MLModel.describe вызывает super().describe(), что для объекта '
     'HybridAnalyzer приводит к EconomicIndicator.describe, затем к '
     'BaseComponent.describe. Результат: «Компонент 3 / '
     'EconomicIndicator / MLModel / HybridAnalyzer». Это демонстрирует, '
     'что super() следует MRO фактического объекта и не означает '
     'просто вызов текстуально указанного родителя.')
text('Если разные родители реализуют validate или log независимо, '
     'победит первый подходящий метод по MRO; важная проверка может '
     'быть пропущена. Я вынесла общие операции в единые миксины. '
     'При необходимости можно переопределить конфликтующий метод, '
     'явно вызвать нужный класс либо изменить порядок родителей. '
     'Для кооперативной работы все участники цепочки должны сохранять '
     'совместимые сигнатуры, передавать оставшиеся **kwargs и '
     'использовать super(). Прямые вызовы двух родительских '
     'конструкторов могут повторно инициализировать общего предка.')
text('Несогласуемый порядок родителей Python отклоняет с TypeError '
     'уже при создании класса: например, X(A, B) и Y(B, A) невозможно '
     'объединить в Z(X, Y). Этот случай отдельно проверен тестом. '
     'Перегрузка арифметических операторов в данной работе не '
     'требуется. В переопределённых методах сохранены совместимые '
     'сигнатуры; входные данные проверяются по типу и диапазону.')
figure(5, 'Создание компонентов, фактический MRO и работа super()', height=16)

heading('2.4. Контрольный пример и UML-моделирование')
text('Я загрузила x = [1, 2, 3, 4, 5] и '
     'y = [100, 200, 300, 400, 500]. LOO MSE для k = 1 равна '
     '10000; для k = 2 — 9000; для k = 3 — 18666,67. '
     'Выбрано k = 2. На обучающих данных MAE = 50 и RMSE = 50. '
     'Для x = 6 ближайшие значения y равны 500 и 400: '
     'прогноз (500 + 400) / 2 = 450. При затратах 300 получен '
     'ROI = (450 − 300) / 300 · 100 % = 50 % (см. рисунок 6). '
     'Повторный прогноз 450 берётся из кэша.')
text('Результат экспортирован методом export_report в файл '
     'example_report.json, включённый в комплект работы. '
     'В демонстрации main JSON выводится в консоль; при генерации '
     'иллюстраций тем же методом записывается отдельный файл. '
     'Идентификатор зависит от порядка создания объектов и '
     'не влияет на вычисления.')
figure(6, 'Подбор k, метрики, прогноз и JSON-отчёт', height=17)
text('Я построила UML-диаграммы с именами классов, полями, методами, '
     'типами и видимостью: «+» — открытый член, «−» — внутренний '
     'член по соглашению Python. Полные описания миксинов приведены '
     'на рисунке 7, основных классов — на рисунке 8. На рисунке 8 '
     'показана основная иерархия с общим предком; все прямые связи '
     'с миксинами отдельно показаны на рисунке 9. Пустой '
     'треугольник направлен к родителю. Композиция классов в этом '
     'ядре не используется.')
figure(7, 'UML-описания миксинов аналитического ядра')
figure(8, 'UML-описания основных классов и общего предка', height=20)
figure(9, 'Полная схема прямых связей наследования аналитического ядра')
text('Диаграмма последовательности показывает создание объекта, '
     'кооперативную инициализацию, загрузку и проверку данных, '
     'подбор k с обучением, получение прогноза, расчёт ROI '
     'и экспорт в JSON-файл (см. рисунок 10). Вызовы родителей '
     'объединены в одну линию жизни для читаемости.')
figure(10, 'Диаграмма последовательности аналитического сценария', height=19)
text('В журнале видны операции load data, fit, tune, cache miss, '
     'cache hit, calculate ROI и export JSON. При нулевых затратах '
     'возбуждается ZeroDivisionError, который перехватывается '
     'в демонстрации (см. рисунок 11). Валидация также отклоняет '
     'строки вместо чисел, bool, NaN, бесконечность, '
     'неравные длины наборов, недопустимые k и отрицательные '
     'затраты. Прогноз до обучения приводит к RuntimeError.')
figure(11, 'Журналирование и обработка деления на ноль', height=14)
listing(3, 'analytics_core.py')

heading('Проверка результатов')
text('Я выполнила 14 автоматических тестов unittest — все завершились '
     'успешно. Проверены смешанная группа сотрудников, уникальность id, '
     'явное разрешение конфликта work и изменение приоритета родителей; '
     'проходы Dense и Conv; фактический MRO и однократная инициализация '
     'BaseComponent; прогноз kNN и разрешение равных расстояний; '
     'подбор k по LOO без включения проверяемой точки; MAE и RMSE; '
     'срабатывание и очистка кэша; ROI и его интерпретация; '
     'полный сценарий с записью и чтением JSON; проверки типов, '
     'диапазонов и несовместимый MRO. Все три программы также '
     'запущены отдельно как консольные приложения. Полные протоколы '
     'запусков и результаты проверок приложены к исходникам.')
text('Для воспроизведения из каталога lab3_shevchenko: '
     'python3 -m unittest discover -s programs -v. '
     'Программы используют только стандартную библиотеку Python; '
     'установка ML-фреймворков не требуется.')
gap('conclusion_before')
p = text('Вывод: я освоила принципы множественного наследования в Python, '
         'изучила механизм MRO и работу super(). Реализовала классы '
         'сотрудников и слоёв нейросети для второго варианта, '
         'разработала аналитическое ядро с миксинами проверки данных, '
         'кэширования, журналирования и экспорта. Выполнила обучение '
         'kNN-регрессора, подбор гиперпараметра, оценку ошибок и '
         'экономический расчёт. Научилась разрешать конфликты '
         'одноимённых методов и предотвращать повторную инициализацию '
         'общего предка. Построила UML-диаграммы классов и '
         'последовательности, проверила решения контрольными '
         'примерами и автоматическими тестами.', kind='conclusion')

used = set()
for element in doc._element.iter():
    for attr in (qn('r:embed'), qn('r:link'), qn('r:id')):
        value = element.get(attr)
        if value:
            used.add(value)
for rid, relation in list(doc.part.rels.items()):
    if relation.reltype.endswith('/image') and rid not in used:
        del doc.part.rels[rid]
doc.core_properties.author = 'Шевченко О.Н.'
doc.core_properties.last_modified_by = 'Шевченко О.Н.'
doc.core_properties.title = 'Лабораторная работа № 3. Вариант 2'
doc.core_properties.subject = 'Множественное наследование'
doc.core_properties.comments = ''
output = ROOT / 'ЛР3_Шевченко_ОН_вариант_2.docx'
doc.save(output)
print(output)
